import json, re, unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
TS=ROOT/'src/lib/official-squads-2026-27.ts'
XLSX=ROOT/'outputs/jogadores-alcunhas-2026-09-05/Jogadores_Nomes_Alcunhas_Girabola.xlsx'
OUT=ROOT/'outputs/auditoria-plantel-alcunhas-2026-09-07'

def norm(v):
    s=unicodedata.normalize('NFD',str(v or ''))
    return re.sub(r'[^a-z0-9]+',' ',''.join(c for c in s if unicodedata.category(c)!='Mn').lower()).strip()

text=TS.read_text()
official=json.loads(text[text.index('['):text.rindex('] as const;')+1])
club_alias={norm(s['club']):s['teamId'] for s in official}
club_alias.update({norm(s['sourceName']):s['teamId'] for s in official})

wb=load_workbook(XLSX,data_only=True,read_only=True)
ws=wb['Jogadores']
headers=[str(v or '').strip() for v in next(ws.iter_rows(min_row=4,max_row=4,values_only=True))]
rows=[]
for vals in ws.iter_rows(min_row=5,values_only=True):
    if not any(v is not None for v in vals): continue
    r=dict(zip(headers,vals)); r['excelRow']=len(rows)+5
    r['teamId']=club_alias.get(norm(r['Clube']))
    rows.append(r)

off=[]
for squad in official:
    for p in squad['players']:
        off.append({**p,'teamId':squad['teamId'],'club':squad['club']})

issues=[]
def add(severity,kind,evidence,impact,fix,confidence='high'):
    issues.append(dict(severity=severity,kind=kind,evidence=evidence,impact=impact,fix=fix,confidence=confidence))

for field in ('maId','fifaId'):
    groups=defaultdict(list)
    for p in off:
        if p[field]: groups[p[field]].append(p)
    for key,items in groups.items():
        if len(items)>1:
            add('high',f'duplicate_official_{field}',{'value':key,'players':[{'teamId':p['teamId'],'fullName':p['fullName']} for p in items]},'O mesmo identificador liga inscrições potencialmente distintas.','Confirmar a inscrição válida e tornar o identificador único.')

for team,items in __import__('itertools').groupby(sorted(off,key=lambda x:x['teamId']),key=lambda x:x['teamId']):
    items=list(items)
    for field in ('jerseyNumber','fullName'):
        groups=defaultdict(list)
        for p in items:
            key=norm(p[field])
            if key: groups[key].append(p)
        for key,ps in groups.items():
            if len(ps)>1:
                add('high' if field=='jerseyNumber' else 'medium',f'duplicate_official_team_{field}',{'teamId':team,'value':ps[0][field],'players':[p['fullName'] for p in ps]},'A ligação por clube e camisola/nome torna-se ambígua.','Confirmar e corrigir o registo oficial.')

unknown=[{'row':r['excelRow'],'club':r['Clube']} for r in rows if not r['teamId']]
if unknown:add('high','unknown_club_in_nicknames',unknown,'Estas linhas não podem ser associadas a um plantel.','Corrigir o nome do clube na folha.')

for field in ('ID jogador',):
    groups=defaultdict(list)
    for r in rows:
        if r[field]:groups[norm(r[field])].append(r)
    for key,rs in groups.items():
        if len(rs)>1:add('high','duplicate_nickname_player_id',{'value':key,'rows':[r['excelRow'] for r in rs],'names':[r['Nome completo atual'] for r in rs]},'Duas linhas podem atualizar o mesmo jogador.','Manter uma linha por ID e rever as restantes.')

for team,rs in __import__('itertools').groupby(sorted([r for r in rows if r['teamId']],key=lambda x:x['teamId']),key=lambda x:x['teamId']):
    rs=list(rs)
    for field in ('Nº camisola','Nome completo atual'):
        groups=defaultdict(list)
        for r in rs:
            k=norm(r[field])
            if k:groups[k].append(r)
        for key,xs in groups.items():
            if len(xs)>1:add('high' if field=='Nº camisola' else 'medium',f'duplicate_nickname_team_{field}',{'teamId':team,'value':xs[0][field],'rows':[x['excelRow'] for x in xs],'names':[x['Nome completo atual'] for x in xs]},'A correspondência dentro do clube fica ambígua.','Rever duplicados e jogadores antigos.')

off_by=defaultdict(list); x_by=defaultdict(list)
for p in off:off_by[p['teamId']].append(p)
for r in rows:
    if r['teamId']:x_by[r['teamId']].append(r)

unmatched_off=[]; unmatched_x=[]; matched=[]
for team in sorted(off_by):
    available=set(range(len(x_by[team])))
    for p in off_by[team]:
        exact=[i for i in available if norm(x_by[team][i]['Nome completo atual'])==norm(p['fullName'])]
        if not exact:
            exact=[i for i in available if norm(x_by[team][i]['Nome completo atual'])==norm(p['name'])]
        if len(exact)==1:
            i=exact[0];available.remove(i);matched.append((p,x_by[team][i]))
        else: unmatched_off.append(p)
    unmatched_x.extend(x_by[team][i] for i in available)

for p in unmatched_off:
    candidates=[r for r in unmatched_x if r['teamId']==p['teamId']]
    scored=sorted(((SequenceMatcher(None,norm(p['fullName']),norm(r['Nome completo atual'])).ratio(),r) for r in candidates),reverse=True,key=lambda x:x[0])
    best=scored[0] if scored else (0,None)
    add('medium','official_player_without_exact_nickname_match',{'teamId':p['teamId'],'officialFullName':p['fullName'],'maId':p['maId'],'bestWorkbookMatch':None if not best[1] else {'row':best[1]['excelRow'],'name':best[1]['Nome completo atual'],'similarity':round(best[0],3)}},'A alcunha ou o nome de exibição pode ficar ligado ao jogador errado ou ausente.','Confirmar pelo ID MA ou FIFA e corrigir o nome.', 'medium')
for r in unmatched_x:
    add('medium','nickname_row_without_exact_official_match',{'teamId':r['teamId'],'row':r['excelRow'],'name':r['Nome completo atual'],'playerId':r['ID jogador']},'A linha pode ser antiga, duplicada ou ter grafia incompatível com o plantel atual.','Confirmar se o jogador continua inscrito e ligar pelo ID MA/FIFA.', 'medium')

for p,r in matched:
    current=norm(r['Alcunha atual']); registered=norm(p['popularName'])
    if registered and current!=registered:
        add('medium','nickname_differs_from_official_popular_name',{'teamId':p['teamId'],'row':r['excelRow'],'fullName':p['fullName'],'officialPopularName':p['popularName'],'workbookNickname':r['Alcunha atual']},'O portal pode mostrar uma alcunha diferente da inscrição.','Escolher a fonte autorizada e uniformizar a alcunha.')
    def shirt(v):
        try: return str(int(str(v).strip()))
        except (ValueError,TypeError): return str(v or '').strip()
    if shirt(r['Nº camisola'])!=shirt(p['jerseyNumber']):
        add('medium','shirt_number_mismatch',{'teamId':p['teamId'],'row':r['excelRow'],'fullName':p['fullName'],'official':p['jerseyNumber'],'workbook':r['Nº camisola']},'Escalações podem associar o jogador errado quando usam a camisola.','Atualizar a camisola usando a inscrição válida.')

summary={'officialPlayers':len(off),'nicknameRows':len(rows),'exactMatches':len(matched),'officialWithoutExactMatch':len(unmatched_off),'nicknameRowsWithoutExactMatch':len(unmatched_x),'issueCount':len(issues),'bySeverity':dict(Counter(i['severity'] for i in issues)),'byKind':dict(Counter(i['kind'] for i in issues))}
OUT.mkdir(parents=True,exist_ok=True)
(OUT/'audit.json').write_text(json.dumps({'summary':summary,'issues':issues},ensure_ascii=False,indent=2)+'\n')
lines=['# Auditoria: Plantel × Alcunhas','',f"Plantel oficial: **{len(off)}** jogadores. Folha de alcunhas: **{len(rows)}** linhas. Correspondências exatas por clube e nome: **{len(matched)}**.",'', '## Resumo','', '| Gravidade | Quantidade |','|---|---:|']
for k in ('high','medium','low'):lines.append(f"| {k} | {summary['bySeverity'].get(k,0)} |")
lines += ['', '## Tipos de inconsistência','', '| Tipo | Quantidade |','|---|---:|']
for k,v in sorted(summary['byKind'].items(),key=lambda x:(-x[1],x[0])):lines.append(f'| {k} | {v} |')
lines += ['', 'Os casos completos, com linhas da folha e sugestões de correção, estão em `audit.json`.','']
(OUT/'README.md').write_text('\n'.join(lines))
print(json.dumps(summary,ensure_ascii=False,indent=2))
